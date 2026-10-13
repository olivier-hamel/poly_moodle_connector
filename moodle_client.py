"""Read-only Moodle API client for authorized personal use."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
import os
import requests

BASE_URL = "https://moodle.polymtl.ca"
TOKEN_FILE = Path(__file__).with_name(".moodle_token")
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 50_000


class MoodleError(RuntimeError):
    """Avoid copying secret-bearing responses into errors."""


def load_token() -> str:
    token = os.environ.get("MOODLE_TOKEN", "").strip()
    if not token and TOKEN_FILE.exists():
        token = TOKEN_FILE.read_text(encoding="utf-8").strip()
    if not token:
        raise MoodleError("No token found. Run: python setup_browser_access.py")
    return token


@dataclass
class MoodleClient:
    token: str = field(repr=False)

    def call(self, function: str, **parameters):
        form = {"wstoken": self.token, "wsfunction": function,
                "moodlewsrestformat": "json", **parameters}
        try:
            response = requests.post(
                f"{BASE_URL}/webservice/rest/server.php",
                data=form, timeout=30, allow_redirects=False,
            )
            if 300 <= response.status_code < 400:
                raise MoodleError("Moodle API redirected; the redirect was not followed.")
            response.raise_for_status()
            data = response.json()
        except (requests.RequestException, ValueError) as exc:
            raise MoodleError("Moodle API request failed (network or invalid response).") from exc
        if isinstance(data, dict) and data.get("exception"):
            # Do not echo server messages that might contain sensitive details.
            raise MoodleError(f"Moodle API error: {data.get('errorcode', 'unknown')}")
        return data

    def site_info(self):
        return self.call("core_webservice_get_site_info")

    def list_courses(self):
        user_id = self.site_info().get("userid")
        if user_id is None:
            raise MoodleError("Moodle did not return your user ID.")
        courses = self.call("core_enrol_get_users_courses", userid=user_id)
        return [
            {"id": c["id"], "shortname": c.get("shortname", ""),
             "fullname": c.get("fullname", "")}
            for c in courses
        ]

    def _course_files(self, course_id: int):
        sections = self.call("core_course_get_contents", courseid=course_id)
        files = []
        for section in sections:
            for module in section.get("modules", []):
                for index, content in enumerate(module.get("contents", [])):
                    if content.get("type") != "file" or not content.get("fileurl"):
                        continue
                    files.append({
                        "id": f"{module['id']}:{index}",
                        "filename": content.get("filename", "unnamed"),
                        "mimetype": content.get("mimetype", ""),
                        "filesize": content.get("filesize", 0),
                        "section": section.get("name", ""),
                        "module": module.get("name", ""),
                        "_url": content["fileurl"],
                    })
        return files

    def list_files(self, course_id: int):
        return [{k: v for k, v in f.items() if k != "_url"}
                for f in self._course_files(course_id)]

    def _download(self, file_url: str) -> bytes:
        parts = urlsplit(file_url)
        base = urlsplit(BASE_URL)
        if (parts.scheme != "https" or parts.netloc != base.netloc
                or parts.username is not None or parts.password is not None):
            raise MoodleError("Moodle provided a file link from an unexpected origin.")
        if parts.path.startswith("/webservice/pluginfile.php/"):
            path = parts.path
        elif parts.path.startswith("/pluginfile.php/"):
            path = "/webservice" + parts.path
        else:
            raise MoodleError("This Moodle file URL is not a supported pluginfile URL.")
        existing_params = [(k, v) for k, v in parse_qsl(parts.query, keep_blank_values=True)
                           if k.lower() != "token"]
        url = urlunsplit(("https", base.netloc, path, urlencode(existing_params), ""))
        try:
            with requests.get(url, params={"token": self.token},
                              stream=True, allow_redirects=False, timeout=45) as response:
                if 300 <= response.status_code < 400:
                    raise MoodleError("Moodle file URL redirected. Refusing to forward token.")
                response.raise_for_status()
                declared = int(response.headers.get("content-length", "0") or 0)
                if declared > MAX_FILE_BYTES:
                    raise MoodleError("File is larger than the 25 MB safety limit.")
                body = bytearray()
                for chunk in response.iter_content(chunk_size=64 * 1024):
                    body.extend(chunk)
                    if len(body) > MAX_FILE_BYTES:
                        raise MoodleError("File is larger than the 25 MB safety limit.")
        except requests.RequestException as exc:
            raise MoodleError("Unable to download file from Moodle.") from exc
        return bytes(body)

    def read_file(self, course_id: int, file_id: str):
        file_item = next((f for f in self._course_files(course_id)
                          if f["id"] == file_id), None)
        if file_item is None:
            raise MoodleError("File not found in that course. Run list_files again.")
        content = self._download(file_item["_url"])
        filename = file_item["filename"]
        suffix = Path(filename).suffix.lower()
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document
            doc = Document(BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                paragraphs.extend(" | ".join(cell.text for cell in row.cells)
                                  for row in table.rows)
            text = "\n".join(paragraphs)
        elif suffix == ".pptx":
            from pptx import Presentation
            ppt = Presentation(BytesIO(content))
            slides = []
            for index, slide in enumerate(ppt.slides, 1):
                shapes = [shape.text for shape in slide.shapes if shape.has_text_frame]
                slides.append(f"Slide {index}\n" + "\n".join(shapes))
            text = "\n\n".join(slides)
        elif suffix in (".txt", ".md", ".csv", ".json", ".html", ".htm"):
            text = content.decode("utf-8-sig", errors="replace")
            if suffix in (".html", ".htm"):
                from bs4 import BeautifulSoup
                text = BeautifulSoup(text, "html.parser").get_text("\n", strip=True)
        else:
            raise MoodleError(f"Reading {suffix or 'this format'} is not implemented yet.")
        return {
            "filename": filename,
            "course_id": course_id,
            "text": text[:MAX_TEXT_CHARS],
            "truncated": len(text) > MAX_TEXT_CHARS,
            "characters_extracted": len(text),
            "note": "Scanned pages or images may require OCR, not included in this starter.",
        }
